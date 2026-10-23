"""Generate small demo documents used by the EduMate acceptance flow."""

from __future__ import annotations

from pathlib import Path

from docx import Document

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SAMPLES = PROJECT_ROOT / "samples"


def build_machine_learning_docx(path: Path) -> None:
    document = Document()
    document.add_heading("机器学习基础讲义", level=0)
    document.add_heading("第一章 机器学习概述", level=1)
    document.add_paragraph(
        "机器学习研究如何让计算机从数据中自动学习规律。按训练数据是否带有标签，"
        "可以分为监督学习、无监督学习和强化学习。"
    )
    document.add_heading("第二章 监督学习", level=1)
    document.add_paragraph(
        "监督学习使用带标签的数据训练模型。常见任务包括分类与回归。"
        "线性回归通过拟合输入特征与连续目标之间的线性关系进行预测；"
        "逻辑回归虽然名称中有回归，但常用于二分类任务。"
    )
    document.add_paragraph(
        "模型在训练集上的误差称为训练误差，在未见数据上的误差称为泛化误差。"
        "当模型过度拟合训练数据时，会出现过拟合；当模型过于简单时，会出现欠拟合。"
    )
    document.add_heading("第三章 模型评估", level=1)
    document.add_paragraph(
        "常用评估方法包括留出法、交叉验证和自助法。分类任务常用准确率、精确率、"
        "召回率和 F1 值；回归任务常用均方误差和平均绝对误差。"
    )
    document.add_paragraph(
        "学习率是优化算法的重要超参数。学习率过大时，损失可能在最优点附近震荡甚至发散；"
        "学习率过小时，收敛速度会很慢。"
    )
    document.save(path)


def build_data_structure_docx(path: Path) -> None:
    document = Document()
    document.add_heading("数据结构与算法基础", level=0)
    document.add_heading("第一章 线性表", level=1)
    document.add_paragraph(
        "线性表是由相同类型元素构成的有序序列。顺序表使用连续内存存储元素，"
        "支持随机访问；链表通过指针连接结点，插入和删除时不需要移动大量元素。"
    )
    document.add_paragraph(
        "单链表插入操作需要先找到前驱结点，时间复杂度为 O(n)；"
        "在已知前驱结点的情况下，插入本身的时间复杂度为 O(1)。"
    )
    document.add_heading("第二章 树与图", level=1)
    document.add_paragraph(
        "二叉树是每个结点最多有两个子结点的树结构。二叉搜索树的中序遍历结果有序，"
        "平均查找时间复杂度为 O(log n)，但在退化为链表时会变成 O(n)。"
    )
    document.save(path)


def main() -> None:
    SAMPLES.mkdir(parents=True, exist_ok=True)
    build_machine_learning_docx(SAMPLES / "机器学习基础.docx")
    build_data_structure_docx(SAMPLES / "数据结构与算法基础.docx")
    print(f"Generated sample documents in {SAMPLES}")


if __name__ == "__main__":
    main()
